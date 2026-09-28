from __future__ import annotations

from io import BytesIO

from docx import Document

from app.services.document_extractors.base import (

    DocumentSection,

    ExtractedDocument,

)

def extract_docx(

    content: bytes,

    filename: str,

    mime_type: str,

) -> ExtractedDocument:

    document = Document(BytesIO(content))

    text_parts: list[str] = []

    structure: list[dict] = []

    sections: list[DocumentSection] = []

    paragraph_index = 0

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if not text:

            continue

        paragraph_index += 1

        text_parts.append(text)

        style_name = (

            paragraph.style.name

            if paragraph.style

            else None

        )

        structure.append(

            {

                "type": "paragraph",

                "number": paragraph_index,

                "style": style_name,

                "text_length": len(text),

            }

        )

        sections.append(

            DocumentSection(

                type="paragraph",

                identifier=f"paragraph-{paragraph_index}",

                text=text,

                metadata={

                    "style": style_name,

                },

            )

        )

    for table_index, table in enumerate(

        document.tables,

        start=1,

    ):

        rows = []

        for row in table.rows:

            rows.append(

                [

                    cell.text.strip()

                    for cell in row.cells

                ]

            )

        table_text = "\n".join(

            " | ".join(row)

            for row in rows

        )

        text_parts.append(table_text)

        structure.append(

            {

                "type": "table",

                "number": table_index,

                "rows": len(table.rows),

                "columns": (

                    len(table.columns)

                    if table.rows

                    else 0

                ),

            }

        )

        sections.append(

            DocumentSection(

                type="table",

                identifier=f"table-{table_index}",

                text=table_text,

                metadata={

                    "rows": len(table.rows),

                    "columns": (

                        len(table.columns)

                        if table.rows

                        else 0

                    ),

                },

            )

        )

    return ExtractedDocument(

        filename=filename,

        file_type="docx",

        mime_type=mime_type,

        size_bytes=len(content),

        text="\n".join(text_parts),

        metadata={

            "paragraph_count": len(

                document.paragraphs

            ),

            "table_count": len(

                document.tables

            ),

        },

        structure=structure,

        sections=sections,

    )