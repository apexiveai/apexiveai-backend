from __future__ import annotations

import io

from pptx import Presentation

from app.services.document_extractors.base import (

    DocumentSection,

    ExtractedDocument,

)

def extract_pptx(

    filename: str,

    mime_type: str,

    file_bytes: bytes,

) -> ExtractedDocument:

    presentation = Presentation(

        io.BytesIO(file_bytes)

    )

    sections = []

    text_parts = []

    for slide_index, slide in enumerate(

        presentation.slides,

        start=1,

    ):

        slide_parts = []

        for shape in slide.shapes:

            if not hasattr(shape, "text"):

                continue

            text = shape.text.strip()

            if text:

                slide_parts.append(text)

        slide_text = "\n".join(

            slide_parts

        )

        text_parts.append(slide_text)

        sections.append(

            DocumentSection(

                type="slide",

                identifier=f"slide-{slide_index}",

                text=slide_text,

                metadata={

                    "slide": slide_index,

                },

            )

        )

    full_text = "\n".join(

        text_parts

    )

    return ExtractedDocument(

        filename=filename,

        file_type="pptx",

        mime_type=mime_type,

        size_bytes=len(file_bytes),

        text=full_text,

        slides=len(presentation.slides),

        metadata={

            "format": "PPTX",

        },

        structure=[

            {

                "type": "slide",

                "index": index + 1,

            }

            for index in range(

                len(presentation.slides)

            )

        ],

        sections=sections,

    )