from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.services.document_extractors.base import (

    DocumentSection,

    ExtractedDocument,

)

def extract_xlsx(

    content: bytes,

    filename: str,

    mime_type: str,

) -> ExtractedDocument:

    workbook = load_workbook(

        BytesIO(content),

        read_only=True,

        data_only=False,

    )

    text_parts: list[str] = []

    structure: list[dict] = []

    sections: list[DocumentSection] = []

    for sheet in workbook.worksheets:

        sheet_lines: list[str] = []

        for row in sheet.iter_rows():

            values = []

            for cell in row:

                if cell.value is None:

                    values.append("")

                else:

                    values.append(str(cell.value))

            if any(

                value.strip()

                for value in values

            ):

                sheet_lines.append(

                    " | ".join(values)

                )

        sheet_text = "\n".join(

            sheet_lines

        )

        text_parts.append(sheet_text)

        structure.append(

            {

                "type": "sheet",

                "name": sheet.title,

                "rows": sheet.max_row,

                "columns": sheet.max_column,

                "text_length": len(

                    sheet_text

                ),

            }

        )

        sections.append(

            DocumentSection(

                type="sheet",

                identifier=f"sheet-{sheet.title}",

                text=sheet_text,

                metadata={

                    "name": sheet.title,

                    "rows": sheet.max_row,

                    "columns": sheet.max_column,

                },

            )

        )

    sheet_count = len(

        workbook.worksheets

    )

    workbook.close()

    return ExtractedDocument(

        filename=filename,

        file_type="xlsx",

        mime_type=mime_type,

        size_bytes=len(content),

        text="\n".join(text_parts),

        sheets=sheet_count,

        metadata={

            "sheet_count": sheet_count,

        },

        structure=structure,

        sections=sections,

    )